# -*- coding: utf-8 -*-
"""Prueba automática de TODAS las herramientas de RESG PDF Studio.

Genera archivos de ejemplo (PDF, PNG, DOCX, XLSX, PPTX, HTML) y ejecuta
cada función del backend comprobando que el resultado se crea de verdad.

Uso:  python tools/selftest.py
"""
import os
import shutil
import sys
import tempfile
import traceback

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

import pymupdf  # noqa: E402
from PIL import Image  # noqa: E402

from resg_pdf.backend import tools as T  # noqa: E402

TMP = tempfile.mkdtemp(prefix="resg_selftest_")
OK, FALLOS, OMITIDOS = [], [], []


def _log(nombre, estado, extra=""):
    signo = {"OK": "[ OK ]", "FALLO": "[FALLO]", "OMIT": "[OMIT]"}[estado]
    print(f"{signo} {nombre:<28} {extra}", flush=True)
    (OK if estado == "OK" else FALLOS if estado == "FALLO"
     else OMITIDOS).append(nombre)


def archivo_ok(ruta):
    """Comprueba que la salida existe y, si es PDF, que se puede abrir."""
    if isinstance(ruta, str) and os.path.isdir(ruta):
        return len(os.listdir(ruta)) > 0
    if not (ruta and os.path.isfile(str(ruta))
            and os.path.getsize(str(ruta)) > 0):
        return False
    if str(ruta).lower().endswith(".pdf"):
        try:
            doc = pymupdf.open(str(ruta))
            paginas = doc.page_count
            doc.close()
            return paginas > 0
        except Exception as e:  # noqa: BLE001
            print(f"      -> PDF ilegible: {e}")
            return False
    return True


def ejecutar(nombre, funcion, files, out, opts):
    try:
        res = funcion(list(files), out, dict(opts), lambda p, m="": None)
        if archivo_ok(res):
            _log(nombre, "OK", f"(-> {os.path.basename(str(res))})")
        else:
            _log(nombre, "FALLO", "no se creó el archivo de salida")
    except Exception as e:  # noqa: BLE001
        _log(nombre, "FALLO", f"{type(e).__name__}: {e}")
        if os.environ.get("RESG_DEBUG"):
            traceback.print_exc()


def ejecutar_error_esperado(nombre, funcion, files, out, opts, falta):
    """Herramientas que necesitan software externo: aviso claro."""
    try:
        res = funcion(list(files), out, dict(opts), lambda p, m="": None)
        if archivo_ok(res):
            _log(nombre, "OK", "funcionó (software externo disponible)")
        else:
            _log(nombre, "FALLO", "sin salida y sin error")
    except RuntimeError as e:
        _log(nombre, "OMIT", f"requiere {falta}: {str(e)[:60]}")
    except Exception as e:  # noqa: BLE001
        _log(nombre, "FALLO", f"{type(e).__name__}: {e}")


def ejecutar_encoge(nombre, funcion, files, out, opts):
    """Ejecuta una herramienta de compresión y exige que el PDF encoja."""
    try:
        res = funcion(list(files), out, dict(opts), lambda p, m="": None)
        if not archivo_ok(res):
            _log(nombre, "FALLO", "no se creó el archivo de salida")
            return
        antes = os.path.getsize(str(files[0]))
        despues = os.path.getsize(str(res))
        if despues < antes:
            _log(nombre, "OK", f"({antes // 1024} KB -> {despues // 1024} KB)")
        else:
            _log(nombre, "FALLO",
                 f"no redujo el tamaño ({antes // 1024} KB -> "
                 f"{despues // 1024} KB)")
    except Exception as e:  # noqa: BLE001
        _log(nombre, "FALLO", f"{type(e).__name__}: {e}")
        if os.environ.get("RESG_DEBUG"):
            traceback.print_exc()


def ejecutar_paginas(nombre, funcion, files, out, opts, esperadas):
    """Ejecuta la herramienta y comprueba que el PDF resultante tenga N páginas."""
    try:
        res = funcion(list(files), out, dict(opts), lambda p, m="": None)
        if not archivo_ok(res):
            _log(nombre, "FALLO", "no se creó el archivo de salida")
            return
        documento = pymupdf.open(str(res))
        paginas = documento.page_count
        documento.close()
        if paginas == esperadas:
            _log(nombre, "OK", f"{paginas} páginas")
        else:
            _log(nombre, "FALLO",
                 f"{paginas} páginas (se esperaban {esperadas})")
    except Exception as e:  # noqa: BLE001
        _log(nombre, "FALLO", f"{type(e).__name__}: {e}")


# ------------------------------------------------------------ archivos base
def crear_muestras():
    m = {}

    # PDF con texto, título y tabla simple
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 80), "TITULO PRINCIPAL DEL DOCUMENTO",
                         fontsize=24)
        page.insert_text((72, 120),
                         f"Este es el contenido de la pagina {i + 1}. "
                         "Dato CONFIDENCIAL de prueba.", fontsize=12)
        page.insert_text((72, 150), "Fila A   Fila B   Fila C", fontsize=12)
        page.insert_text((72, 170), "1        2        3", fontsize=12)
    m["pdf"] = os.path.join(TMP, "muestra.pdf")
    doc.save(m["pdf"])
    doc.close()

    # Segundo PDF (ligeramente distinto) para comparar y unir
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 80), "TITULO PRINCIPAL DEL DOCUMENTO",
                         fontsize=24)
        page.insert_text((72, 120),
                         f"Este es el contenido de la pagina {i + 1}. "
                         "Dato diferente de prueba.", fontsize=12)
    m["pdf2"] = os.path.join(TMP, "muestra2.pdf")
    doc.save(m["pdf2"])
    doc.close()

    # PDF con formulario interactivo
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Formulario", fontsize=18)
    for j, nombre in enumerate(("nombre", "apellido")):
        w = pymupdf.Widget()
        w.field_name = nombre
        w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        w.rect = pymupdf.Rect(72, 100 + j * 40, 320, 124 + j * 40)
        w.field_value = ""
        page.add_widget(w)
    m["form"] = os.path.join(TMP, "formulario.pdf")
    doc.save(m["form"])
    doc.close()

    # Imagen PNG (rectángulo azul)
    img = Image.new("RGBA", (400, 200), (0, 0, 0, 0))
    for x in range(50, 350):
        for y in range(40, 160):
            img.putpixel((x, y), (40, 90, 200, 255))
    m["png"] = os.path.join(TMP, "imagen.png")
    img.save(m["png"])

    # Firma PNG
    firma = Image.new("RGBA", (300, 120), (0, 0, 0, 0))
    for x in range(20, 280):
        firma.putpixel((x, 60), (10, 10, 10, 255))
        firma.putpixel((x, 61), (10, 10, 10, 255))
    m["firma"] = os.path.join(TMP, "firma.png")
    firma.save(m["firma"])

    # HTML
    m["html"] = os.path.join(TMP, "pagina.html")
    with open(m["html"], "w", encoding="utf-8") as fh:
        fh.write("<html><body><h1>Hola</h1><p>Página de prueba.</p>"
                 "<table border='1'><tr><td>A</td><td>B</td></tr></table>"
                 "</body></html>")

    # Foto grande (para recompresión de imágenes y renderizado)
    try:
        import numpy as np

        alto, ancho = 1400, 1000
        arr = np.zeros((alto, ancho, 3), dtype=np.uint8)
        arr[:, :, 0] = np.arange(ancho, dtype=np.uint8)[None, :]
        arr[:, :, 1] = np.arange(alto, dtype=np.uint8)[:, None]
        arr[:, :, 2] = (np.arange(ancho, dtype=np.uint8)[None, :] // 3 + 90)
        foto = Image.fromarray(arr)
    except Exception:                       # sin numpy: ruido aleatorio
        foto = Image.effect_noise((1000, 1400), 90).convert("RGB")
    m["foto"] = os.path.join(TMP, "foto.jpg")
    foto.save(m["foto"], quality=97)

    # PDF con imágenes (para comprobar compresión y calidad de imagen)
    doc = pymupdf.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_image(pymupdf.Rect(60, 60, 535, 400), filename=m["foto"])
        page.insert_text((72, 440), f"Pagina con foto {i + 1}", fontsize=14)
    m["pdf_img"] = os.path.join(TMP, "con_imagenes.pdf")
    doc.save(m["pdf_img"])
    doc.close()

    # PDF con otro tamaño de página (para el comparador)
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=595)          # A5
    page.insert_text((60, 90), "TITULO PRINCIPAL DEL DOCUMENTO", fontsize=18)
    page.insert_text((60, 130), "Contenido en un tamano distinto.",
                     fontsize=11)
    m["pdf_a5"] = os.path.join(TMP, "a5.pdf")
    doc.save(m["pdf_a5"])
    doc.close()

    # DOCX / XLSX / PPTX
    try:
        import docx
        d = docx.Document()
        d.add_heading("Documento Word de prueba", 0)
        d.add_paragraph("Un párrafo con texto.")
        m["docx"] = os.path.join(TMP, "doc.docx")
        d.save(m["docx"])
    except Exception:
        m["docx"] = None

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["A", "B", "C"])
        ws.append([1, 2, 3])
        m["xlsx"] = os.path.join(TMP, "hoja.xlsx")
        wb.save(m["xlsx"])
    except Exception:
        m["xlsx"] = None

    try:
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Diapositiva de prueba"
        m["pptx"] = os.path.join(TMP, "pres.pptx")
        prs.save(m["pptx"])
    except Exception:
        m["pptx"] = None

    return m


def _p(*a):
    return os.path.join(TMP, *a)


def main():
    m = crear_muestras()

    print("\n=== ORGANIZAR ===")
    ejecutar("Unir PDF", T.unir_pdf, [m["pdf"], m["pdf2"]], _p("unido.pdf"),
             {"password": ""})
    ejecutar("Dividir (cada página)", T.dividir_pdf, [m["pdf"]], _p("div"),
             {"modo": "cada_pagina", "rangos": "1-", "cada": 2,
              "password": ""})
    ejecutar("Dividir (rangos)", T.dividir_pdf, [m["pdf"]], _p("div2"),
             {"modo": "rangos", "rangos": "1-2,3", "cada": 2,
              "password": ""})
    ejecutar("Dividir (cada N)", T.dividir_pdf, [m["pdf"]], _p("div3"),
             {"modo": "tamano", "rangos": "1-", "cada": 2, "password": ""})
    ejecutar("Ordenar PDF", T.ordenar_pdf, [m["pdf"]], _p("ordenado.pdf"),
             {"orden": "3,1,2", "password": ""})
    ejecutar("Rotar PDF", T.rotar_pdf, [m["pdf"]], _p("rotado.pdf"),
             {"grados": "90", "paginas": "", "password": ""})

    print("\n=== OPTIMIZAR ===")
    ejecutar("Comprimir PDF", T.comprimir_pdf, [m["pdf"]],
             _p("comprimido.pdf"), {"nivel": "medio", "password": ""})
    ejecutar_encoge("Comprimir con imágenes", T.comprimir_pdf, [m["pdf_img"]],
                    _p("comprimido_img.pdf"), {"nivel": "bajo", "password": ""})
    ejecutar("Reparar PDF", T.reparar_pdf, [m["pdf"]], _p("reparado.pdf"), {})
    ejecutar_error_esperado("OCR PDF", T.ocr_pdf, [m["pdf"]], _p("ocr.pdf"),
                            {"idioma": "spa", "dpi": "200", "password": ""},
                            "Tesseract OCR")

    print("\n=== CONVERTIR DESDE PDF ===")
    ejecutar("PDF a Word", T.pdf_a_word, [m["pdf"]], _p("salida.docx"), {})
    ejecutar("PDF a PowerPoint", T.pdf_a_ppt, [m["pdf"]], _p("salida.pptx"),
             {"dpi": "150"})
    ejecutar("PDF a Excel", T.pdf_a_excel, [m["pdf"]], _p("salida.xlsx"),
             {"password": ""})
    ejecutar("PDF a JPG", T.pdf_a_jpg, [m["pdf"]], _p("imgs"),
             {"dpi": "150", "formato": "jpg", "password": ""})
    ejecutar("PDF a JPG (con fotos)", T.pdf_a_jpg, [m["pdf_img"]], _p("imgs3"),
             {"dpi": "96", "formato": "jpg", "password": ""})
    ejecutar("PDF a PNG", T.pdf_a_jpg, [m["pdf"]], _p("imgs2"),
             {"dpi": "96", "formato": "png", "password": ""})
    ejecutar("PDF a Markdown", T.pdf_a_markdown, [m["pdf"]], _p("salida.md"),
             {"password": ""})
    ejecutar_error_esperado("PDF a PDF/A", T.pdf_a_pdfa, [m["pdf"]],
                            _p("salida_a.pdf"), {"nivel": "2"},
                            "Ghostscript")

    print("\n=== CONVERTIR A PDF ===")
    ejecutar("JPG a PDF", T.jpg_a_pdf, [m["png"]], _p("desde_img.pdf"),
             {"orientacion": "auto", "margen": "0"})
    ejecutar("HTML a PDF", T.html_a_pdf, [m["html"]], _p("desde_html.pdf"),
             {"url": ""})
    ejecutar("Escanea a PDF (imagen)", T.escanear_a_pdf, [m["png"]],
             _p("escaneado.pdf"), {})
    # --- Office -> PDF con el motor interno (sin LibreOffice) ---
    soffice_real = T.buscar_soffice
    T.buscar_soffice = lambda: None
    try:
        if m["docx"]:
            ejecutar("Word a PDF (motor interno)", T.office_a_pdf, [m["docx"]],
                     _p("doc_interno.pdf"), {})
        if m["pptx"]:
            ejecutar("PowerPoint a PDF (motor interno)", T.office_a_pdf,
                     [m["pptx"]], _p("pres_interno.pdf"), {})
        if m["xlsx"]:
            ejecutar("Excel a PDF (motor interno)", T.office_a_pdf,
                     [m["xlsx"]], _p("hoja_interno.pdf"), {})
    finally:
        T.buscar_soffice = soffice_real
    if T.buscar_soffice():
        if m["docx"]:
            ejecutar("Word a PDF (LibreOffice)", T.office_a_pdf, [m["docx"]],
                     _p("doc.pdf"), {})
        if m["pptx"]:
            ejecutar("PowerPoint a PDF (LibreOffice)", T.office_a_pdf,
                     [m["pptx"]], _p("pres.pdf"), {})
        if m["xlsx"]:
            ejecutar("Excel a PDF (LibreOffice)", T.office_a_pdf, [m["xlsx"]],
                     _p("hoja.pdf"), {})
    else:
        print("[INFO] LibreOffice no instalado: se omite la prueba de alta "
              "fidelidad (el motor interno ya se validó arriba).")


    print("\n=== EDITAR ===")
    ejecutar("Editar PDF", T.editar_pdf, [m["pdf"]], _p("editado.pdf"),
             {"texto": "Hola mundo", "pagina": 1, "x": 10, "y": 30,
              "tamano": 14, "color": "rojo", "password": ""})
    ejecutar("Marca de agua (texto)", T.marca_agua_pdf, [m["pdf"]],
             _p("marca.pdf"),
             {"tipo": "texto", "texto": "CONFIDENCIAL", "imagen": "",
              "posicion": "centro", "rotacion": 45, "opacidad": 30,
              "tamano": 48, "paginas": "", "password": ""})
    ejecutar("Marca de agua (imagen)", T.marca_agua_pdf, [m["pdf"]],
             _p("marca_img.pdf"),
             {"tipo": "imagen", "texto": "", "imagen": m["png"],
              "posicion": "centro", "rotacion": 0, "opacidad": 50,
              "tamano": 30, "paginas": "", "password": ""})
    ejecutar("Firmar PDF", T.firmar_pdf, [m["pdf"]], _p("firmado.pdf"),
             {"imagen": m["firma"], "posicion": "inferior derecha",
              "paginas": "", "tamano": 25, "password": ""})
    ejecutar("Números de página", T.numeros_pagina_pdf, [m["pdf"]],
             _p("numeros.pdf"),
             {"formato": "Página {n} de {t}", "posicion": "inferior centro",
              "tamano": 11, "inicio": 1, "paginas": "", "password": ""})
    ejecutar("Recortar PDF", T.recortar_pdf, [m["pdf"]], _p("recortado.pdf"),
             {"izquierda": 5, "derecha": 5, "superior": 5, "inferior": 5,
              "paginas": "", "password": ""})
    ejecutar("Censurar (texto)", T.censurar_pdf, [m["pdf"]],
             _p("censurado.pdf"),
             {"modo": "texto", "buscar": "CONFIDENCIAL", "pagina": 1,
              "area": "10,10,50,20", "password": ""})
    ejecutar("Censurar (área)", T.censurar_pdf, [m["pdf"]],
             _p("censurado2.pdf"),
             {"modo": "area", "buscar": "", "pagina": 1,
              "area": "10,10,50,20", "password": ""})
    ejecutar("Formularios (listar)", T.formularios_pdf, [m["form"]],
             _p("campos.txt"), {"modo": "listar", "valores": "",
                                "password": ""})
    ejecutar("Formularios (rellenar)", T.formularios_pdf, [m["form"]],
             _p("relleno.pdf"),
             {"modo": "rellenar", "valores": "nombre=Ana\napellido=Perez",
              "password": ""})

    print("\n=== SEGURIDAD ===")
    protegido = _p("protegido.pdf")
    ejecutar("Proteger PDF", T.proteger_pdf, [m["pdf"]], protegido,
             {"clave": "1234", "clave_propietario": "", "password": ""})
    ejecutar("Desbloquear PDF", T.desbloquear_pdf, [protegido],
             _p("desbloqueado.pdf"), {"password": "1234"})
    ejecutar("Comparar PDF", T.comparar_pdf, [m["pdf"], m["pdf2"]],
             _p("comparado.pdf"), {"dpi": "100"})
    ejecutar("Comparar tamaños distintos", T.comparar_pdf,
             [m["pdf"], m["pdf_a5"]], _p("comparado2.pdf"), {"dpi": "72"})
    ejecutar("Comparar idénticos", T.comparar_pdf, [m["pdf"], m["pdf"]],
             _p("comparado3.pdf"), {"dpi": "72"})

    print("\n=== ORGANIZAR CON LA VISTA DE MINIATURAS ===")
    plan = [{"file": m["pdf"], "page": 2, "rot": 90},
            {"file": m["pdf"], "page": 0, "rot": 0},
            {"file": m["png"], "page": 0, "rot": 0}]
    ejecutar_paginas("Plan de páginas (PDF+imagen)", T.aplicar_plan,
                     [m["pdf"], m["png"]], _p("plan.pdf"),
                     {"plan": plan, "password": ""}, 3)
    ejecutar_paginas("Unir con plan de páginas", T.unir_pdf,
                     [m["pdf"], m["pdf2"]], _p("unido_plan.pdf"),
                     {"plan": [{"file": m["pdf"], "page": 0, "rot": 0},
                               {"file": m["pdf2"], "page": 2, "rot": 0}],
                      "password": ""}, 2)
    ejecutar_paginas("Ordenar con plan", T.ordenar_pdf, [m["pdf"]],
                     _p("orden_plan.pdf"),
                     {"plan": [{"file": m["pdf"], "page": i, "rot": 0}
                               for i in (2, 1, 0)], "password": ""}, 3)
    ejecutar_paginas("Rotar con plan", T.rotar_pdf, [m["pdf"]],
                     _p("rotado_plan.pdf"),
                     {"plan": [{"file": m["pdf"], "page": 0, "rot": 90}],
                      "password": ""}, 1)
    ejecutar_paginas("Eliminar páginas (rangos)", T.eliminar_paginas, [m["pdf"]],
                     _p("sin_p2.pdf"), {"paginas": "2", "password": ""}, 2)
    ejecutar_paginas("Eliminar páginas (desmarcadas)", T.eliminar_paginas,
                     [m["pdf"]], _p("sin_marcadas.pdf"),
                     {"plan": [{"file": m["pdf"], "page": 0, "rot": 0},
                               {"file": m["pdf"], "page": 2, "rot": 0}],
                      "password": ""}, 2)
    ejecutar_paginas("Extraer páginas (un PDF)", T.extraer_paginas, [m["pdf"]],
                     _p("extraido.pdf"),
                     {"modo": "unico", "rangos": "1,3", "password": ""}, 2)
    ejecutar("Extraer páginas (uno por rango)", T.extraer_paginas, [m["pdf"]],
             os.path.join(TMP, "extraidos"),
             {"modo": "separados", "rangos": "1-2,3", "password": ""})

    print("\n=== EDITOR VISUAL ===")
    elementos = [
        {"tipo": "texto", "pagina": 1, "x": 8, "y": 8, "w": 60, "h": 10,
         "texto": "Texto del editor", "tamano": 18, "color": "#d02b34",
         "fuente": "helv", "opacidad": 100},
        {"tipo": "rectangulo", "pagina": 1, "x": 8, "y": 25, "w": 40, "h": 8,
         "color": "#ffd633", "opacidad": 80},
        {"tipo": "elipse", "pagina": 2, "x": 45, "y": 40, "w": 25, "h": 20,
         "color": "#8c4ac6", "opacidad": 100},
        {"tipo": "linea", "pagina": 2, "x": 10, "y": 70, "w": 45, "h": 12,
         "color": "rojo", "grosor": 2},
        {"tipo": "resaltado", "pagina": 3, "x": 6, "y": 20, "w": 45, "h": 6,
         "color": "#ffd633", "opacidad": 45},
        {"tipo": "enlace", "pagina": 3, "x": 8, "y": 40, "w": 35, "h": 6,
         "url": "https://resg-tech.example"},
        {"tipo": "imagen", "pagina": 3, "x": 8, "y": 55, "w": 30, "h": 20,
         "imagen": m["png"], "proporcion": True},
    ]
    ejecutar_paginas("Editar PDF (elementos del editor)", T.editar_pdf,
                     [m["pdf"]], _p("editor.pdf"),
                     {"elementos": elementos, "password": ""}, 3)

    print("\n=== INTELIGENCIA DE DOCUMENTO ===")
    largo = os.path.join(TMP, "largo.pdf")
    documento = pymupdf.open()
    for i in range(2):
        pagina = documento.new_page()
        parrafos = " ".join(
            "La gestión documental local permite trabajar sin internet y sin "
            f"cuentas de usuario. Párrafo {n} de la página {i + 1} con datos "
            "suficientes para que el resumen automático tenga de dónde elegir."
            for n in range(10))
        pagina.insert_textbox(pymupdf.Rect(60, 60, 540, 770), parrafos,
                              fontsize=11)
    documento.save(largo)
    documento.close()
    ejecutar("Resumir PDF (Markdown)", T.resumir_pdf, [largo],
             _p("resumen.md"), {"formato": "md", "frases": 6,
                                "password": ""})
    ejecutar("Resumir PDF (texto)", T.resumir_pdf, [largo],
             _p("resumen.txt"), {"formato": "txt", "frases": 5,
                                 "password": ""})
    ejecutar("Resumir PDF (PDF)", T.resumir_pdf, [largo],
             _p("resumen.pdf"), {"formato": "pdf", "frases": 4,
                                 "password": ""})
    if T.idiomas_traduccion():
        ejecutar("Traducir PDF", T.traducir_pdf, [m["pdf"]],
                 _p("traducido.pdf"),
                 {"origen": "es", "destino": "en", "formato": "pdf",
                  "password": ""})
    else:
        ejecutar_error_esperado("Traducir PDF", T.traducir_pdf, [m["pdf"]],
                                _p("traducido.pdf"),
                                {"origen": "es", "destino": "en"},
                                "argostranslate")

    print("\n=== FORMULARIOS (CREAR CAMPOS) ===")
    campos = "\n".join([
        "texto;1;10;10;40;6;nombre;Ana",
        "casilla;1;10;20;6;6;acepto;si",
        "desplegable;1;10;30;40;6;ciudad;Madrid|Bilbao|Sevilla",
        "firma;2;10;40;40;12;firma_cliente;",
    ])
    ejecutar("Formularios (crear campos)", T.formularios_pdf, [m["pdf"]],
             _p("campos_creados.pdf"),
             {"modo": "crear", "campos": campos, "password": ""})

    print("\n" + "=" * 62)
    print(f"OK: {len(OK)}   FALLOS: {len(FALLOS)}   "
          f"OMITIDOS: {len(OMITIDOS)}")
    if FALLOS:
        print("FALLARON: " + ", ".join(FALLOS))
    print(f"Temp de trabajo: {TMP}")
    shutil.rmtree(TMP, ignore_errors=True)
    return 1 if FALLOS else 0


if __name__ == "__main__":
    sys.exit(main())

